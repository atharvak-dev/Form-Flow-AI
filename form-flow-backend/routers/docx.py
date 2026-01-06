"""
Word Document Router

Endpoints for uploading and filling Word documents.
"""

from fastapi import APIRouter, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
from typing import Dict, Any
import io
import uuid

from services.docx.docx_parser import parse_docx_fields, DocxParser
from utils.logging import get_logger

logger = get_logger(__name__)

router = APIRouter(prefix="/docx", tags=["Word Documents"])

# In-memory storage for uploaded documents (replace with Redis/DB for production)
_docx_storage: Dict[str, bytes] = {}


@router.post("/upload")
async def upload_docx(file: UploadFile = File(...)) -> Dict[str, Any]:
    """
    Upload and parse a Word document to extract fillable fields.
    
    Supports:
    - Bracket placeholders: [Name], [Email]
    - Underscore placeholders: ____
    - Content Controls (SDT)
    
    Returns:
        Schema with detected fields
    """
    # Validate file type
    if not file.filename.lower().endswith(('.docx', '.doc')):
        raise HTTPException(
            status_code=400,
            detail="Only .docx files are supported. Please upload a valid Word document."
        )
    
    try:
        # Read file content
        content = await file.read()
        
        if file.filename.lower().endswith('.doc'):
            raise HTTPException(
                status_code=400,
                detail="Legacy .doc format is not supported. Please save as .docx and re-upload."
            )
        
        # Parse document
        result = parse_docx_fields(content)
        
        if not result.get("success"):
            raise HTTPException(
                status_code=400,
                detail=result.get("message", "Failed to parse document")
            )
        
        # Store for later filling
        docx_id = result.get("docx_id")
        _docx_storage[docx_id] = content
        
        logger.info(f"Parsed Word document: {file.filename}, {result.get('total_fields')} fields found")
        
        return result
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error processing Word document: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail=f"Failed to process document: {str(e)}"
        )


@router.post("/fill")
async def fill_docx(docx_id: str, data: Dict[str, str]) -> Dict[str, Any]:
    """
    Fill a previously uploaded Word document with provided data.
    
    This replaces placeholders with actual values.
    """
    if docx_id not in _docx_storage:
        raise HTTPException(
            status_code=404,
            detail="Document not found. Please re-upload."
        )
    
    try:
        from docx import Document
        
        # Load original document
        content = _docx_storage[docx_id]
        doc = Document(io.BytesIO(content))
        
        # Replace placeholders in paragraphs
        for paragraph in doc.paragraphs:
            for field_name, value in data.items():
                # Replace bracket placeholders
                import re
                # Match both [Name] and [name] style
                pattern = re.compile(rf'\[{re.escape(field_name)}\]', re.IGNORECASE)
                if pattern.search(paragraph.text):
                    for run in paragraph.runs:
                        run.text = pattern.sub(value, run.text)
                        
                # Also try display name variations
                display_pattern = re.compile(rf'\[{field_name.replace("_", " ")}\]', re.IGNORECASE)
                if display_pattern.search(paragraph.text):
                    for run in paragraph.runs:
                        run.text = display_pattern.sub(value, run.text)
        
        # Save filled document
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        
        # Generate download ID
        download_id = str(uuid.uuid4())
        _docx_storage[f"filled_{download_id}"] = output.getvalue()
        
        return {
            "success": True,
            "download_id": download_id,
            "fields_filled": len(data),
            "message": "Document filled successfully"
        }
        
    except Exception as e:
        logger.error(f"Error filling document: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail=f"Failed to fill document: {str(e)}"
        )


@router.get("/download/{download_id}")
async def download_docx(download_id: str):
    """Download a filled Word document."""
    storage_key = f"filled_{download_id}"
    
    if storage_key not in _docx_storage:
        raise HTTPException(
            status_code=404,
            detail="Document not found or expired."
        )
    
    content = _docx_storage[storage_key]
    
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f"attachment; filename=filled_document.docx"
        }
    )
